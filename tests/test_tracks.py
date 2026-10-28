"""
The EdTech label on the shared Full Stack Product Engineer build.

The two EdTech developer postings invite to the same assessment as Full Stack
Product Engineer, so their candidates arrive on role 28 looking like everyone
else. These pin the three places the difference has to survive: the lookup that
writes it, the label read back off a submission, and the spreadsheet column a
hiring manager filters on.
"""

import io

import openpyxl
import pytest

from backend.grading import tier_resolver
from backend.mail import shortlist

EDTECH = "CD4D14E6F6"       # Developer (EdTech, AI)
EDTECH_LEAD = "0895DC92C6"  # Lead Developer (EdTech, AI)


class FakeStore:
    def __init__(self, subs):
        self.subs = {s["_id"]: s for s in subs}

    def submissions_missing_track(self, job_id):
        return [s for s in self.subs.values()
                if s["job_id"] == job_id and "posting_track" not in s]

    def set_posting_track(self, submission_id, label, shortcode=None):
        self.subs[submission_id]["posting_track"] = {
            "label": label, "shortcode": shortcode}


ROLE = {"_id": 28, "slug": "product-engineer",
        "title": "Full Stack Product Engineer"}


def test_both_edtech_postings_feed_the_product_engineer_build():
    assert tier_resolver.tracked_postings("product-engineer") == {
        EDTECH: "EdTech", EDTECH_LEAD: "EdTech"}
    assert tier_resolver.tracked_postings("workflow-analyst") == {}


def test_resolve_labels_edtech_by_email_and_defaults_the_rest(monkeypatch):
    lists = {EDTECH: [{"email": "Ed@Example.com"}], EDTECH_LEAD: []}
    monkeypatch.setattr(tier_resolver, "get_job_candidates",
                        lambda shortcode: lists[shortcode])
    store = FakeStore([
        {"_id": 1, "job_id": 28, "candidate_email": "ed@example.com"},
        {"_id": 2, "job_id": 28, "candidate_email": "fs@example.com"},
    ])

    result = tier_resolver.resolve_tracks(ROLE, store)

    assert result["counts"] == {"EdTech": 1, "Full Stack Dev": 1}
    assert store.subs[1]["posting_track"] == {"label": "EdTech",
                                              "shortcode": EDTECH}
    assert store.subs[2]["posting_track"]["label"] == "Full Stack Dev"


def test_a_workable_outage_writes_nothing(monkeypatch):
    # Half a pass would stamp the default on EdTech candidates whose posting
    # was simply the one that failed to load.
    def boom(shortcode):
        if shortcode == EDTECH_LEAD:
            raise RuntimeError("503")
        return []
    monkeypatch.setattr(tier_resolver, "get_job_candidates", boom)
    store = FakeStore([{"_id": 1, "job_id": 28, "candidate_email": "a@b.c"}])

    with pytest.raises(tier_resolver.TierResolutionFailed):
        tier_resolver.resolve_tracks(ROLE, store)
    assert "posting_track" not in store.subs[1]
    assert tier_resolver.ensure_tracks(ROLE, store) is None


def test_label_falls_back_to_the_role_default_and_is_absent_elsewhere():
    edtech = {"job_id": 28, "posting_track": {"label": "EdTech"}}
    unresolved = {"job_id": 28}
    assert tier_resolver.track_label(edtech) == "EdTech"
    assert tier_resolver.track_label(unresolved) == "Full Stack Dev"
    # The dashboard marks EdTech only.
    assert tier_resolver.track_label(edtech, include_default=False) == "EdTech"
    assert tier_resolver.track_label(unresolved, include_default=False) is None
    # Full Stack Developer (portal 17) is a different build with no tracks.
    assert tier_resolver.track_label({"job_id": 17}) is None


def test_spreadsheet_carries_a_track_column_beside_the_name():
    subs = [
        {"_id": 1, "job_id": 28, "candidate_name": "Ed",
         "posting_track": {"label": "EdTech"}},
        {"_id": 2, "job_id": 28, "candidate_name": "Fay"},
    ]
    rows = [shortlist.row(s, rank=i, include_scores=False)
            for i, s in enumerate(subs, start=1)]
    blob = shortlist.build_xlsx(ROLE, rows)

    sheet = openpyxl.load_workbook(io.BytesIO(blob)).active
    headers = [cell.value for cell in sheet[4]]
    assert headers[:3] == ["#", "Candidate", "Track"]
    assert [sheet.cell(row=r, column=3).value for r in (5, 6)] == [
        "EdTech", "Full Stack Dev"]


def test_other_roles_sheets_are_unchanged():
    rows = [shortlist.row({"_id": 1, "job_id": 17, "candidate_name": "A"},
                          rank=1, include_scores=False)]
    assert "track" not in rows[0]
    blob = shortlist.build_xlsx({"title": "Full Stack Developer"}, rows)
    headers = [c.value for c in openpyxl.load_workbook(io.BytesIO(blob)).active[4]]
    assert "Track" not in headers
